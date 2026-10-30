import sys
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
NAVY="0D2240"; ORANGE="FF6A2B"; YEL="FFF6C7"; LINE="E6E1D6"
th=Side(style="thin",color=LINE); B=Border(bottom=th)
USD='"$"#,##0.00'
def build(ws):
    ws.sheet_view.showGridLines=False
    for c,w in zip("ABCDEF",(44,10,10,14,16,4)): ws.column_dimensions[c].width=w
    for r in (1,2,3):
        for c in "ABCDE": ws[f"{c}{r}"].fill=PatternFill("solid",fgColor=NAVY)
    ws["A1"]="[Your Business Name]"; ws["A1"].font=Font(size=18,bold=True,color="FFFFFF")
    ws["A2"]="[Address] · [Phone] · [Email] · License # [000000]"; ws["A2"].font=Font(size=10,color="C3CFDF")
    ws["E1"]="INVOICE"; ws["E1"].font=Font(size=16,bold=True,color=ORANGE); ws["E1"].alignment=Alignment(horizontal="right")
    ws.row_dimensions[1].height=30
    y=PatternFill("solid",fgColor=YEL)
    for r,(l,v) in enumerate([("Invoice #","2001"),("Invoice date","2026-11-20"),("Due date","2026-12-05"),("Estimate #","1001")],5):
        ws[f"D{r}"]=l; ws[f"D{r}"].font=Font(bold=True,color="586273"); ws[f"E{r}"]=v; ws[f"E{r}"].fill=y
    ws["A5"]="Bill to"; ws["A5"].font=Font(bold=True,color="586273")
    for r,v in enumerate(["[Customer name]","[Job address]","[Phone / email]"],6): ws[f"A{r}"]=v; ws[f"A{r}"].fill=y
    ws["A10"]="Work completed"; ws["A10"].font=Font(bold=True,color="586273")
    ws["A11"]="[Describe the finished work and the completion date]"; ws["A11"].fill=y
    for i,h in enumerate(["Item / description","Qty","Unit","Unit price","Amount"]):
        c=ws.cell(13,i+1,h); c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="14325C")
        if i: c.alignment=Alignment(horizontal="right")
    rows=[("Labor",16,"hr",75),("Materials (see list)",1,"lot",1200),("Equipment rental",1,"day",150),
          ("Disposal / dump fee",1,"load",95),("Change order #1 (signed)",1,"ea",822)]
    for r in range(14,26):
        e=rows[r-14] if r-14<len(rows) else None
        for col in range(1,5):
            c=ws.cell(r,col,e[col-1] if e else None); c.fill=y; c.border=B
            if col==4: c.number_format=USD
        a=ws.cell(r,5,f'=IF(OR(B{r}="",D{r}=""),"",B{r}*D{r})'); a.number_format=USD; a.border=B
    def tot(r,l,f,big=False):
        ws[f"D{r}"]=l; ws[f"D{r}"].font=Font(bold=True,color="0F1B2D"); ws[f"E{r}"]=f; ws[f"E{r}"].number_format=USD
        ws[f"E{r}"].font=Font(bold=True,size=14 if big else 11,color=ORANGE if big else "0F1B2D")
        ws[f"D{r}"].alignment=Alignment(horizontal="right")
    ws["A28"]="Taxable amount (e.g. materials) $"; ws["B28"]=1200; ws["B28"].fill=y; ws["B28"].number_format=USD
    ws["A29"]="Sales tax rate %"; ws["B29"]=0; ws["B29"].fill=y
    ws["A30"]="Deposit already received $"; ws["B30"]=793.5; ws["B30"].fill=y; ws["B30"].number_format=USD
    ws["A31"]="Other payments received $"; ws["B31"]=0; ws["B31"].fill=y; ws["B31"].number_format=USD
    tot(27,"Subtotal","=SUM(E14:E25)")
    tot(28,"Sales tax","=B28*B29/100")
    tot(29,"Invoice total","=E27+E28")
    tot(30,"Less deposit received","=-B30")
    tot(31,"Less other payments","=-B31")
    tot(32,"BALANCE DUE","=E29+E30+E31",True)
    ws["A34"]="Payment"; ws["A34"].font=Font(bold=True,color="586273")
    terms=["[How to pay: check payable to ..., bank transfer, card link. Edit to match your business.]",
           "Payment due by the due date above. [Late fee policy, if any, as allowed by your state and your contract.]",
           "Thank you for your business."]
    for i,t in enumerate(terms,35): ws[f"A{i}"]=t; ws[f"A{i}"].font=Font(size=10,color="586273")
    ws["A39"]="Template by Clever Loop Digital · cleverloopdigital.com"; ws["A39"].font=Font(size=9,italic=True,color="9AA6BA")
    ws.print_area="A1:E39"; ws.page_setup.fitToWidth=1
wb=Workbook(); ws=wb.active; ws.title="Invoice"; build(ws)
wb.save("../../public/templates/contractor-invoice-template.xlsx")
p="../../docs/distribution/product/contractor-pricing-toolkit.xlsx"
wb=load_workbook(p)
if "Customer Invoice" not in wb.sheetnames:
    ws=wb.create_sheet("Customer Invoice"); build(ws)
    st=wb["Start Here"]
    for r in range(17,10,-1):
        st[f"A{r+1}"]=st[f"A{r}"].value; st[f"A{r+1}"].font=st[f"A{r}"].font.copy()
    st["A11"]="7. Customer Invoice: bill the finished job, subtract the deposit and show the balance due."
    st["A11"].font=Font(size=11,color="0F1B2D")
    wb.save(p)
print(wb.sheetnames)
