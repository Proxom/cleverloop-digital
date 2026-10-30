import sys
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
NAVY="0D2240"; ORANGE="FF6A2B"; YEL="FFF6C7"; SOFT="F7F5F0"; LINE="E6E1D6"
th=Side(style="thin",color=LINE); B=Border(bottom=th)
USD='"$"#,##0.00'
def build(ws):
    ws.sheet_view.showGridLines=False
    for c,w in zip("ABCDEF",(44,10,10,14,16,4)): ws.column_dimensions[c].width=w
    for r in (1,2,3):
        for c in "ABCDE": ws[f"{c}{r}"].fill=PatternFill("solid",fgColor=NAVY)
    ws["A1"]="[Your Business Name]"; ws["A1"].font=Font(size=18,bold=True,color="FFFFFF")
    ws["A2"]="[Address] · [Phone] · [Email] · License # [000000]"; ws["A2"].font=Font(size=10,color="C3CFDF")
    ws["E1"]="ESTIMATE"; ws["E1"].font=Font(size=16,bold=True,color=ORANGE); ws["E1"].alignment=Alignment(horizontal="right")
    ws.row_dimensions[1].height=30
    y=PatternFill("solid",fgColor=YEL)
    for r,(l,v) in enumerate([("Estimate #","1001"),("Date","2026-11-02"),("Valid until","2026-12-02")],5):
        ws[f"D{r}"]=l; ws[f"D{r}"].font=Font(bold=True,color="586273"); ws[f"E{r}"]=v; ws[f"E{r}"].fill=y
    ws["A5"]="Prepared for"; ws["A5"].font=Font(bold=True,color="586273")
    for r,v in enumerate(["[Customer name]","[Job address]","[Phone / email]"],6): ws[f"A{r}"]=v; ws[f"A{r}"].fill=y
    ws["A10"]="Job description"; ws["A10"].font=Font(bold=True,color="586273")
    ws["A11"]="[Describe the work in one or two sentences]"; ws["A11"].fill=y
    hdr=["Item / description","Qty","Unit","Unit price","Amount"]
    for i,h in enumerate(hdr):
        c=ws.cell(13,i+1,h); c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="14325C")
        if i: c.alignment=Alignment(horizontal="right")
    rows=[("Labor",16,"hr",75),("Materials (see list)",1,"lot",1200),("Equipment rental",1,"day",150),("Disposal / dump fee",1,"load",95)]
    for r in range(14,26):
        e=rows[r-14] if r-14<len(rows) else None
        for col in range(1,5):
            c=ws.cell(r,col,e[col-1] if e else None); c.fill=y; c.border=B
            if col==4: c.number_format=USD
        a=ws.cell(r,5,f'=IF(OR(B{r}="",D{r}=""),"",B{r}*D{r})'); a.number_format=USD; a.border=B
    def tot(r,l,f,big=False):
        ws[f"D{r}"]=l; ws[f"D{r}"].font=Font(bold=True,color="0F1B2D"); ws[f"E{r}"]=f; ws[f"E{r}"].number_format=USD
        ws[f"E{r}"].font=Font(bold=True,size=14 if big else 11,color=ORANGE if big else "0F1B2D")
        ws[f"D{r}"].alignment=Alignment(horizontal="right")
    tot(27,"Subtotal","=SUM(E14:E25)")
    ws["A28"]="Taxable amount (e.g. materials) $"; ws["B28"]=1200; ws["B28"].fill=y; ws["B28"].number_format=USD
    ws["A29"]="Sales tax rate %"; ws["B29"]=0; ws["B29"].fill=y
    ws["A30"]="Deposit %"; ws["B30"]=30; ws["B30"].fill=y
    tot(28,"Sales tax","=B28*B29/100")
    tot(29,"TOTAL","=E27+E28",True)
    tot(30,"Deposit due to start","=E29*B30/100")
    tot(31,"Balance on completion","=E29-E30")
    ws["A33"]="Terms"; ws["A33"].font=Font(bold=True,color="586273")
    terms=["[Edit to match your business and your state's rules.]",
           "Price valid until the date above. Work not listed here is extra and needs a signed change order.",
           "Deposit due on acceptance; balance due on completion. [Accepted payment methods.]"]
    for i,t in enumerate(terms,34): ws[f"A{i}"]=t; ws[f"A{i}"].font=Font(size=10,color="586273")
    ws["A38"]="Accepted by (customer signature): ______________________     Date: __________"
    ws["A40"]="Template by Clever Loop Digital · cleverloopdigital.com"; ws["A40"].font=Font(size=9,italic=True,color="9AA6BA")
    ws.print_area="A1:E40"; ws.page_setup.fitToWidth=1
# free standalone
wb=Workbook(); ws=wb.active; ws.title="Estimate"; build(ws)
wb.save("../../public/templates/contractor-estimate-template.xlsx")
# add to paid toolkit
p="../../docs/distribution/product/contractor-pricing-toolkit.xlsx"
wb=load_workbook(p)
if "Customer Estimate" not in wb.sheetnames:
    ws=wb.create_sheet("Customer Estimate"); build(ws)
    st=wb["Start Here"]
    for r in range(17,4,-1):
        v=st[f"A{r}"].value; st[f"A{r+1}"]=v; st[f"A{r+1}"].font=st[f"A{r}"].font.copy()
    st["A9"]="6. Customer Estimate: a clean estimate to send the customer, with line items, tax, deposit and signature line."
    st["A9"].font=Font(size=11,color="0F1B2D")
    wb.save(p)
print(wb.sheetnames)
