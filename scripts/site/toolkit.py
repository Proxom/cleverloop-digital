from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.worksheet.datavalidation import DataValidation
NAVY="0D2240"; ORANGE="FF6A2B"; YEL="FFF6C7"; SOFT="F7F5F0"; LINE="E6E1D6"
th=Side(style="thin",color=LINE); B=Border(bottom=th)
USD='"$"#,##0.00'; PCT='0.0"%"'
wb=Workbook()
def sheet(ws,title,sub,width=(46,18,4,40)):
    ws.sheet_view.showGridLines=False
    for i,w in enumerate(width): ws.column_dimensions["ABCD"[i]].width=w
    ws["A1"]=title; ws["A1"].font=Font(size=18,bold=True,color="FFFFFF")
    ws["A2"]=sub; ws["A2"].font=Font(size=10,color="C3CFDF")
    for r in (1,2):
        for c in "ABCD": ws[f"{c}{r}"].fill=PatternFill("solid",fgColor=NAVY)
    ws.row_dimensions[1].height=32
    ws.protection.sheet=False
def head(ws,r,t):
    ws[f"A{r}"]=t; ws[f"A{r}"].font=Font(bold=True,size=9,color="586273")
    for c in "AB": ws[f"{c}{r}"].fill=PatternFill("solid",fgColor=SOFT)
def inp(ws,r,label,val,fmt=None,note=""):
    ws[f"A{r}"]=label; c=ws[f"B{r}"]; c.value=val; c.fill=PatternFill("solid",fgColor=YEL); c.font=Font(bold=True,color="4A3B00")
    if fmt: c.number_format=fmt
    ws[f"A{r}"].border=B; c.border=B
    if note: ws[f"D{r}"]=note; ws[f"D{r}"].font=Font(size=9,italic=True,color="586273")
def out(ws,r,label,f,fmt=USD,big=False,note=""):
    ws[f"A{r}"]=label; c=ws[f"B{r}"]; c.value=f; c.number_format=fmt
    c.font=Font(bold=True,size=14 if big else 11,color=ORANGE if big else "0F1B2D")
    if big:
        for x in "AB": ws[f"{x}{r}"].fill=PatternFill("solid",fgColor=NAVY)
        ws[f"A{r}"].font=Font(bold=True,color="FFFFFF",size=12)
    ws[f"A{r}"].border=B; c.border=B
    if note: ws[f"D{r}"]=note; ws[f"D{r}"].font=Font(size=9,italic=True,color="586273")

# 0 Start
ws=wb.active; ws.title="Start Here"
sheet(ws,"Contractor Pricing Toolkit","Clever Loop Digital · cleverloopdigital.com",(100,4,4,4))
lines=["How to use this workbook",
"1. Type only in the YELLOW cells. Everything else is a formula.",
"2. Quote & Margin: price one job from materials, labor, subs, overhead and your target margin.",
"3. Hourly Rate: the rate you must charge to reach your income goal after overhead and unbillable time.",
"4. Change Order: price extra work added mid-job (labor + disruption allowance + marked-up materials).",
"5. Job Profit Log: list finished jobs and see the margin you really made on each, and in total.",
"",
"Margin = profit ÷ price.  Markup = profit ÷ cost.  A 30% margin needs a 42.9% markup.",
"Google Sheets: File > Import > Upload this file. Excel: just open it.",
"",
"The numbers already filled in are examples. Replace them with your own.",
"Figures are planning estimates only. This workbook is not accounting, tax or legal advice.",
"Questions: hello@cleverloopdigital.com"]
for i,t in enumerate(lines,4):
    ws[f"A{i}"]=t; ws[f"A{i}"].font=Font(bold=(i==4),size=13 if i==4 else 11,color="0F1B2D")

# 1 Quote & Margin
ws=wb.create_sheet("Quote & Margin"); sheet(ws,"Quote & Margin","Price one job at the margin you want")
head(ws,4,"INPUTS")
inp(ws,5,"Materials cost ($)",1200,USD)
inp(ws,6,"Labor hours",24,'0.00')
inp(ws,7,"Labor cost per hour ($) - wages + payroll costs",35,USD,"Use the Hourly Rate tab or your payroll records")
inp(ws,8,"Subcontractor / rental cost ($)",300,USD)
inp(ws,9,"Overhead (% of direct cost)",12,PCT,"Yearly overhead ÷ yearly direct job costs")
inp(ws,10,"Target margin (% of price)",30,PCT)
inp(ws,11,"Sales tax on materials (%) - 0 if not charged",0,PCT,"Check your state's rules")
inp(ws,12,"Deposit requested up front (%)",30,PCT)
head(ws,14,"RESULTS")
out(ws,15,"Labor total","=B6*B7")
out(ws,16,"Direct cost","=B5+B15+B8")
out(ws,17,"Overhead","=B16*B9/100")
out(ws,18,"Total cost (break-even price)","=B16+B17")
out(ws,19,"QUOTE PRICE (before tax)","=IF(B10>=100,0,B18/(1-B10/100))",big=True)
out(ws,20,"Sales tax","=B5*B11/100")
out(ws,21,"Quote price incl. tax","=B19+B20")
out(ws,22,"Profit on this job","=B19-B18")
out(ws,23,"Markup on cost","=IF(B18=0,0,B22/B18*100)",PCT)
out(ws,24,"Profit per labor hour","=IF(B6=0,0,B22/B6)")
out(ws,25,"Deposit due","=B21*B12/100")
head(ws,27,"WHAT IF")
out(ws,28,"Margin if client negotiates 10% off","=IF(B19=0,0,(B19*0.9-B18)/(B19*0.9)*100)",PCT)
out(ws,29,"Margin if client negotiates 20% off","=IF(B19=0,0,(B19*0.8-B18)/(B19*0.8)*100)",PCT)

# 2 Hourly rate
ws=wb.create_sheet("Hourly Rate"); sheet(ws,"Hourly Rate","The rate you need to hit your income goal")
head(ws,4,"INPUTS")
inp(ws,5,"Income you want to pay yourself per year ($)",60000,USD)
inp(ws,6,"Yearly overhead: insurance, truck, tools, software, phone ($)",18000,USD)
inp(ws,7,"Working weeks per year",48,'0')
inp(ws,8,"Hours worked per week",45,'0')
inp(ws,9,"Billable share of hours (%)",65,PCT,"Driving, quoting and admin are not billable")
inp(ws,10,"Profit margin on top (%)",15,PCT)
head(ws,12,"RESULTS")
out(ws,13,"Hours worked per year","=B7*B8",'#,##0')
out(ws,14,"Billable hours per year","=B13*B9/100",'#,##0')
out(ws,15,"Cost to cover per year","=B5+B6")
out(ws,16,"Break-even hourly rate","=IF(B14=0,0,B15/B14)")
out(ws,17,"HOURLY RATE TO CHARGE","=IF(B10>=100,0,B16/(1-B10/100))",big=True)

# 3 Change order
ws=wb.create_sheet("Change Order"); sheet(ws,"Change Order","Price extra work added mid-job, then get it signed")
head(ws,4,"INPUTS")
inp(ws,5,"Extra labor hours",6,'0.00')
inp(ws,6,"Billable rate per hour ($)",70,USD)
inp(ws,7,"Disruption allowance (% of labor)",10,PCT,"Stop/restart time, small orders, resequencing")
inp(ws,8,"Extra materials at cost ($)",300,USD)
inp(ws,9,"Materials markup (%)",20,PCT)
head(ws,11,"RESULTS")
out(ws,12,"Labor","=B5*B6")
out(ws,13,"Disruption","=B12*B7/100")
out(ws,14,"Materials charged","=B8*(1+B9/100)")
out(ws,15,"CHANGE ORDER PRICE","=B12+B13+B14",big=True)
ws["A17"]="Write down the scope, price, schedule change and payment timing, and get it signed before starting."
ws["A17"].font=Font(italic=True,size=10,color="586273")

# 4 Job log
ws=wb.create_sheet("Job Profit Log"); ws.sheet_view.showGridLines=False
hdr=["Date","Customer / job","Quoted price ($)","Materials ($)","Labor ($)","Subs & other ($)","Overhead %","Total cost ($)","Profit ($)","Margin %"]
for i,w in enumerate([12,30,16,14,12,16,12,15,13,11]): ws.column_dimensions[chr(65+i)].width=w
ws["A1"]="Job Profit Log"; ws["A1"].font=Font(size=18,bold=True,color="FFFFFF")
ws["A2"]="What you really made on each job"; ws["A2"].font=Font(size=10,color="C3CFDF")
for r in (1,2):
    for i in range(10): ws.cell(r,i+1).fill=PatternFill("solid",fgColor=NAVY)
for i,h in enumerate(hdr):
    c=ws.cell(4,i+1,h); c.font=Font(bold=True,color="FFFFFF",size=10); c.fill=PatternFill("solid",fgColor="14325C"); c.alignment=Alignment(wrap_text=True,vertical="center")
ex=[("2026-10-01","Example: bathroom repaint",2400,380,900,0,12),("2026-10-08","Example: deck repair",3744,1200,840,300,12)]
for r in range(5,35):
    e=ex[r-5] if r-5<len(ex) else None
    for col in range(1,8):
        c=ws.cell(r,col, e[col-1] if e else None); c.fill=PatternFill("solid",fgColor=YEL); c.border=B
        if col in(3,4,5,6): c.number_format=USD
        if col==7: c.number_format=PCT
    ws.cell(r,8,f'=IF(C{r}="","",(D{r}+E{r}+F{r})*(1+G{r}/100))').number_format=USD
    ws.cell(r,9,f'=IF(C{r}="","",C{r}-H{r})').number_format=USD
    ws.cell(r,10,f'=IF(OR(C{r}="",C{r}=0),"",I{r}/C{r}*100)').number_format=PCT
    for col in (8,9,10): ws.cell(r,col).border=B
ws["B36"]="TOTAL"; ws["B36"].font=Font(bold=True)
ws["C36"]="=SUM(C5:C34)"; ws["H36"]="=SUM(H5:H34)"; ws["I36"]="=SUM(I5:I34)"; ws["J36"]='=IF(C36=0,"",I36/C36*100)'
for c,f in (("C36",USD),("H36",USD),("I36",USD),("J36",PCT)): ws[c].number_format=f; ws[c].font=Font(bold=True,color=ORANGE,size=12)
ws.freeze_panes="A5"
wb.save("../../docs/distribution/product/contractor-pricing-toolkit.xlsx")
print("saved")
